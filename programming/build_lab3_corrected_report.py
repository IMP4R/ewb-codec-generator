"""Generate Lab3 report from authoritative page9 and real live verification."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
import json,shutil
ROOT=Path('/Users/user0/АВУ/сдача/Лаб_3');REV=Path('/Users/user0/ewb_rev')
p=ROOT/'Отчёт_Лаб_3.docx'
b=REV/'programming/original_reports/Отчёт_Лаб_3.before_round6.docx'
b.parent.mkdir(parents=True,exist_ok=True)
if p.exists() and not b.exists():shutil.copy2(p,b)
d=Document();sec=d.sections[0];sec.left_margin=Cm(3);sec.right_margin=Cm(1.5)
d.styles['Normal'].font.name='Times New Roman';d.styles['Normal'].font.size=Pt(12)
d.add_paragraph('ВЯТСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ')
d.add_heading('Лабораторная работа №3',0)
d.add_paragraph('Логические вентили и синтез комбинационных схем\nВариант 1\nВыполнил: Денис Лузянин\nКиров, 2026');d.add_page_break()
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=str(s)
d.add_heading('1. Транзисторные вентили',1)
d.add_paragraph('НЕ выполнен на одном NPN; И-НЕ — на двух последовательных NPN; ИЛИ-НЕ — на двух параллельных NPN. Резистор нагрузки 1 кОм, питание 5 В. Inverter.ewb сохранён без изменения. Схемы NAND и NOR открыты в EWB по рисункам методички.')
for name,shot in [('НЕ','Inverter_verified.png'),('И-НЕ','NAND_exact.png'),('ИЛИ-НЕ','NOR_exact.png')]:
 d.add_heading(name,2);d.add_picture(str(REV/'screenshots'/shot),width=Cm(16))
d.add_heading('Таблицы напряжений: статус измерений',2)
d.add_paragraph('Входные уровни по методичке: 0,1 В и 2 В. Численные выходные напряжения необходимо сверить в работающей EWB. Значения 0,18/0,35 В из прежнего отчёта не подтверждены текущими скриншотами и исключены. Следующая таблица содержит ожидаемые логические уровни, а не измеренные напряжения.')
table(['U1, В','U2, В','И-НЕ, лог.','ИЛИ-НЕ, лог.'],[(.1,.1,1,1),(.1,2,1,0),(2,.1,1,0),(2,2,0,0)])
d.add_heading('2. Таблица варианта 1',1)
d.add_paragraph('Источник: Лабораторная работа 3.pdf, страница 9, таблица «Вариант 1». Таблица примера на странице 6 не является заданием варианта 1.')
vector=[1,0,1,1,0,0,0,1]
table(['x','y','z','f'],[(i>>2&1,i>>1&1,i&1,f) for i,f in enumerate(vector)])
d.add_heading('3. СДНФ и упрощение',1)
d.add_paragraph('f = ¬x·¬y·¬z + ¬x·y·¬z + ¬x·y·z + x·y·z.')
d.add_paragraph('Группируем первый и второй термы, третий и четвёртый: f = ¬x·¬z·(¬y+y) + y·z·(¬x+x) = ¬x·¬z + y·z.')
d.add_paragraph('Сокращённая КНФ: f = (¬x+z)(y+¬z). Раскрытие даёт ¬xy + ¬x¬z + yz; терм ¬xy — консенсусный, поэтому избыточен.')
d.add_heading('4. Базисы И-НЕ и ИЛИ-НЕ',1)
d.add_paragraph('И-НЕ: f = ¬(¬(¬x·¬z) · ¬(y·z)). Два входных инвертора и три двухвходовых И-НЕ.')
d.add_paragraph('ИЛИ-НЕ: f = ¬(¬(¬x+z) + ¬(y+¬z)). Два входных инвертора и три двухвходовых ИЛИ-НЕ.')
d.add_heading('5. СКНФ',1)
d.add_paragraph('Нулевые наборы: 001,100,101,110. f = (x+y+¬z)(¬x+y+z)(¬x+y+¬z)(¬x+¬y+z). Использованы три НЕ, четыре ИЛИ3 и один И4.')
d.add_heading('6. Число транзисторов',1)
d.add_paragraph('Считаем в учебной модели на NPN: НЕ — 1; И-НЕ_n и ИЛИ-НЕ_n — n; И_n и ИЛИ_n — n+1. Это оценка эквивалентной реализации вентилей, не физический состав идеальных DIGITAL-компонентов EWB. Питание, переключатели и индикаторы в подсчёт не входят.')
table(['Схема','Состав','Транзисторы'],[('СДНФ','3 НЕ + 4 И3 + ИЛИ4','3+4·4+5=24'),('Упрощённая','2 НЕ + 2 И2 + ИЛИ2','2+2·3+3=11'),('И-НЕ','2 НЕ + 3 И-НЕ2','2+3·2=8'),('ИЛИ-НЕ','2 НЕ + 3 ИЛИ-НЕ2','2+3·2=8'),('СКНФ','3 НЕ + 4 ИЛИ3 + И4','3+4·4+5=24')])
d.add_heading('7. Проверка в EWB',1)
for title,folder in [('СДНФ','sdnf'),('СКНФ','sknf'),('И/ИЛИ/НЕ','min_andornot'),('И-НЕ','min_nand'),('ИЛИ-НЕ','min_nor')]:
 directory=REV/'screenshots'/(folder+'_verified');rows=json.loads((directory/'results.json').read_text())
 assert len(rows)==8 and all(r['observed_output']==r['expected'] for r in rows)
 d.add_heading(title,2);d.add_paragraph('Восемь входных наборов проверены на работающей модели; входы и выход считаны по красным индикаторам с реальных скриншотов.')
 table(['xyz','EWB f','Требуется'],[(r['inputs'],r['observed_output'],r['expected']) for r in rows])
 d.add_picture(str(directory/'111.png'),width=Cm(16))
d.add_heading('Вывод и оставшиеся действия',1)
d.add_paragraph('Пять функциональных схем соответствуют таблице варианта 1; 40/40 live-комбинаций совпали. Упрощение снижает учебную оценку с 24 до 11 транзисторов, реализации И-НЕ и ИЛИ-НЕ — до 8. Для окончательной сдачи требуется добавить подтверждённые таблицы выходных напряжений транзисторных вентилей; данный отчёт не утверждает, что эти измерения завершены.')
d.save(p);Document(p);print(p)
