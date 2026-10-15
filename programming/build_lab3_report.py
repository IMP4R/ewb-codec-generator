from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
import json,shutil
ROOT=Path('/Users/user0/АВУ/сдача/Лаб_3');REV=Path('/Users/user0/ewb_rev');d=Document();d.styles['Normal'].font.name='Times New Roman';d.styles['Normal'].font.size=Pt(14)
s=d.sections[0];s.left_margin=Cm(3);s.right_margin=Cm(1.5)
d.add_paragraph('ВЯТСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ');d.add_heading('ОТЧЁТ',0);d.add_paragraph('Лабораторная работа №3\nАрхитектура вычислительных устройств\nЛогические вентили и синтез комбинационных схем\nВариант 1');d.add_paragraph('Выполнил: Денис Лузянин\nКиров, 2026');d.add_page_break()
d.add_heading('1. Транзисторные логические вентили',1);d.add_paragraph('В EWB построены инвертор, NAND с последовательными NPN-транзисторами и NOR с параллельными NPN-транзисторами. Питание 5 В, коллекторная нагрузка 1 кОм. Низкий уровень входа 100 мВ, высокий — 2 В. Схемы используют штатные символы DIN.')
measure=json.loads((REV/'codec_v2/measurements/actual_voltages.json').read_text())['cases']
for gate,figure in [('Inverter','Inverter_verified.png'),('NAND','NAND_final.png'),('NOR','NOR_fixed.png')]:
 d.add_heading(gate,2);d.add_picture(str(REV/'screenshots'/figure),width=Cm(16));t=d.add_table(rows=1,cols=3);t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,['Входные напряжения, В','Выход, В','Логический выход']):c.text=v
 for r in measure:
  if r['gate']!=gate:continue
  row=t.add_row().cells;row[0].text='; '.join(str(x) for x in r['inputs_V']);row[1].text=str(r['output_V']);row[2].text=str(r['logic'])
 d.add_paragraph('Измерения выполнены в настоящем EWB 5.12. Полные скриншоты каждого запуска сохранены в /Users/user0/ewb_rev/screenshots/; имена файлов записаны в actual_voltages.json.')
d.add_heading('2. Функция варианта 1',1);t=d.add_table(rows=1,cols=4);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['x','y','z','f']):c.text=v
for i,f in enumerate([1,0,1,1,0,0,0,1]):
 for c,v in zip(t.add_row().cells,[(i>>2)&1,(i>>1)&1,i&1,f]):c.text=str(v)
d.add_paragraph('СДНФ: f = ¬x¬y¬z ∨ ¬xy¬z ∨ ¬xyz ∨ xyz.\nУпрощение: первые два слагаемых дают ¬x¬z; последние два дают yz. Итого f = ¬x¬z ∨ yz.\nСКНФ: (x∨y∨¬z)∧(¬x∨y∨z)∧(¬x∨y∨¬z)∧(¬x∨¬y∨z).\nУпрощённая КНФ: (¬x∨z)∧(y∨¬z).')
d.add_paragraph('Базис NAND: f = NAND(NAND(¬x,¬z), NAND(y,z)).\nБазис NOR: f = NOR(NOR(¬x,z), NOR(y,¬z)). Инверсии реализуются инверторами.')
d.add_heading('3. Проверка цифровых реализаций',1);d.add_paragraph('Пять схем реализованы на штатных DIGITAL-вентилях EWB. Независимый проверяющий скрипт восстанавливает сети по номерам физических выводов компонентов и проверяет все восемь комбинаций каждой реализации. Это проверка топологии и логической модели, отдельно от измерений аналоговых вентилей в EWB. Результат всех пяти таблиц совпадает с вариантом 1.')
d.add_heading('4. Сравнение числа транзисторов',1);d.add_paragraph('Учебная оценка для транзисторной реализации: НЕ — 1 транзистор; NAND(n), NOR(n) — n; AND(n), OR(n) — n+1. Библиотечные идеальные DIGITAL-компоненты не содержат физической модели транзисторов, поэтому ниже приведена именно оценка эквивалентной реализации, а не измеренное число элементов внутри модели.')
t=d.add_table(rows=1,cols=3);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Реализация','Расчёт','Транзисторов']):c.text=v
for vals in [('СДНФ','3 НЕ + 4 AND3 + 3 OR2',28),('Минимальная AND/OR/NOT','2 НЕ + 2 AND2 + OR2',11),('NAND','2 НЕ + 3 NAND2',8),('NOR','2 НЕ + 3 NOR2',8),('СКНФ','3 НЕ + 4 OR3 + 3 AND2',28)]:
 for c,v in zip(t.add_row().cells,vals):c.text=str(v)
d.add_heading('Вывод',1);d.add_paragraph('Измеренные напряжения подтверждают инверсию, NAND и NOR. Минимизация функции сокращает количество вентилей и оценочное число транзисторов. Все упрощённые базисы эквивалентны заданной таблице.')
p=ROOT/'Отчёт_Лаб_3.docx';backup=REV/'programming/original_reports/Лаб_3_Отчёт_Лаб_3.docx';backup.parent.mkdir(exist_ok=True)
if p.exists() and not backup.exists():backup.write_bytes(p.read_bytes())
d.save(p);Document(p);print(p)
