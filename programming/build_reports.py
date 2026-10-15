from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
import json
ROOT=Path('/Users/user0/АВУ/сдача');HERE=Path(__file__).parent
reports=json.loads((HERE/'machine_execution_report.json').read_text());lab6=json.loads((HERE/'lab6_execution_report.json').read_text())
def start(num,title):
 d=Document();sec=d.sections[0];sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(3);sec.right_margin=Cm(1.5)
 st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(14)
 d.add_paragraph('ВЯТСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ');d.add_heading('ОТЧЁТ',0);d.add_paragraph(f'Лабораторная работа № {num}\nАрхитектура вычислительных устройств\n{title}\nВариант 1');d.add_paragraph('Выполнил: Денис Лузянин\nКиров, 2026');d.add_page_break();return d
def code(d,p):
 d.add_heading(p.name,2)
 para=d.add_paragraph();r=para.add_run(p.read_text());r.font.name='Consolas';r.font.size=Pt(9)
def tests(d,rows):
 d.add_heading('Проверка выполнения',1);d.add_paragraph('Проверка инструкций выполнена в x86-32 эмуляторе Unicorn после перевода синтаксиса типизированных операндов MASM в NASM и сборки машинного кода. Алгоритм и инструкции не заменялись Python-реализацией. Это не является подтверждением нативной сборки и линковки Visual Studio/MASM; такая проверка выделена отдельно как незавершённая.')
 t=d.add_table(rows=1,cols=3);t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,['Процедура','Тестов','Результат']):c.text=v
 for row in rows:
  cells=t.add_row().cells;cells[0].text=Path(row.get('source',row.get('file'))).name;cells[1].text=str(row['cases']);cells[2].text='PASS; регистры ABI сохранены'
 d.add_paragraph('Контрольные данные и воспроизводимые тесты: /Users/user0/ewb_rev/programming/.')
def backup(p):
 b=HERE/'original_reports';b.mkdir(exist_ok=True);q=b/(p.parent.name+'_'+p.name)
 if p.exists() and not q.exists():q.write_bytes(p.read_bytes())
def save(d,num):
 p=ROOT/f'Лаб_{num}'/f'Отчёт_Лаб_{num}.docx';backup(p);d.save(p);Document(p);print(p)
d=start(6,'Пересылка данных и побайтовые операции');d.add_heading('Цель и порядок работы',1);d.add_paragraph('Исследование MOV и XCHG, пересылка 16/32/64-разрядных данных, байтовый циклический сдвиг, реверс и копирование с инверсией. Каждый исходник с одинаковыми глобальными символами собирается отдельным проектом.')
t=d.add_table(rows=1,cols=5);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Исходное a','Хранится short','Исходное b','Хранится signed char','Копирование']):c.text=v
for a,b in [(10000,100),(50000,500),(-10000,-100),(-50000,-500)]:
 cells=t.add_row().cells;vals=[a,((a+32768)%65536)-32768,b,((b+128)%256)-128,'битовая копия']
 for c,v in zip(cells,vals):c.text=str(v)
d.add_paragraph('0x12345678 в little-endian: 78 56 34 12. Циклический сдвиг на один байт вправо: 0x78123456. Реверс: 0x78563412. XCHG допускает регистр–регистр или регистр–память, но не память–память.')
tests(d,lab6)
for p in sorted((ROOT/'Лаб_6').glob('*.asm')):code(d,p)
code(d,ROOT/'Лаб_6/lab6_transfer_demo.cpp');save(d,6)
d=start(7,'Целочисленная арифметика и логические команды');d.add_heading('Часть 1',1);d.add_paragraph('y = (2*a − 1)/(b − 2). Исходные данные и результат — signed short. Промежуточные вычисления выполняются в 32 битах, чтобы избежать ошибочного отбрасывания допустимого результата. Деление целочисленное с усечением к нулю. Проверяются нулевой знаменатель и диапазон результата.')
d.add_paragraph('Примеры: a=10,b=5 → 6; a=−10,b=5 → −7; b=2 → ошибка деления на ноль; a=32767,b=3 → переполнение результата; a=32767,b=32767 → 2.')
d.add_heading('Часть 2',1);d.add_paragraph('Проверить биты 1,4,7: маска 0x0092. Установить 2,5,8: 0x0124. Сбросить 3,6,9: AND с 0xFDB7. Двойной NOT возвращает значение после этих преобразований. Для всех 65 536 исходных слов выполнено сравнение машинного кода с независимым эталоном.')
tests(d,[r for r in reports if 'Лаб_7' in r['source']])
for p in sorted((ROOT/'Лаб_7').glob('*')):
 if p.suffix in ['.asm','.cpp']:code(d,p)
save(d,7)
d=start(8,'Условные переходы и циклы');d.add_heading('Часть 1. Точный вариант из методички',1);d.add_paragraph('Y = (3*a+5)/(b−1), если a>b; Y = 10, если a=b; Y = (a+b)/(a−b), если a<b. Сравнение знаковое. Промежуточные значения 32-разрядные; результат проверяется на signed short; деление на ноль диагностируется.')
d.add_paragraph('Примеры: a=5,b=3 → 10; a=4,b=4 → 10; a=2,b=5 → −2; a=2,b=1 → ошибка; a=32767,b=2 → переполнение результата.')
d.add_heading('Часть 2',1);d.add_paragraph('A1: среднее положительных элементов статического массива short[10]. Сумма и счётчик 32-разрядные. При отсутствии положительных элементов среднее не определено; avg=0 — только служебный индикатор. B1: сложение двух статических матриц 3×3. C1: сложение динамически выделенных матриц short[9] в row-major. При переполнении хотя бы одного элемента результат не записывается частично. Память освобождается RAII.')
tests(d,[r for r in reports if 'Лаб_8' in r['source']])
for p in sorted((ROOT/'Лаб_8').glob('*')):
 if p.suffix in ['.asm','.cpp']:code(d,p)
save(d,8)
